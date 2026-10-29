"""Run with:  python tests/test_project.py   (or: pytest tests)"""
import glob
import os
import shutil
import sys
import tempfile

import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tools"))

from compatibility_db import (PTF_FILL, VER, CompatibilityDB, build_excel, extract_ptf_ids, needs_ptf,  # noqa: E402
                              ptf_worklist, unique_ptf_ids)
import import_xlsx  # noqa: E402

REF = os.path.join(ROOT, "tests", "reference_zOS_3_2_filtered.xlsx")
DB = CompatibilityDB.load(os.path.join(ROOT, "data"))


def test_loads_all_rows():
    assert len(DB.df) == 240
    assert set(DB.df.vendor) == {"ASG", "BMC", "Broadcom", "IBM", "Precisely", "Rocket"}


def test_extract_ids():
    assert extract_ptf_ids("Yes - LU18313 (corrects PE PTF LU18249); LU19547 if IGGDAU01 has RMID UJ98216") == ["LU18313", "LU19547"]
    assert extract_ptf_ids("also LU05157/58/59 prerequisites") == ["LU05157", "LU05158", "LU05159"]
    assert extract_ptf_ids("Yes - ESP 12.0.08 maintenance level") == []
    assert extract_ptf_ids("No") == []
    assert extract_ptf_ids("Yes - RO78780, LU09856, LU09856") == ["RO78780", "LU09856"]


def test_four_rows_that_used_to_say_no_now_need_ptfs():
    d = DB.df
    for prod, vrm in [("Compress Data Compression", "5.5"), ("Earl", "6.1"),
                      ("ESP Workload Automation", "12.0"), ("IDMS/DB", "19.0")]:
        row = d[(d["product"] == prod) & (d["vrm"] == vrm)].iloc[0]
        assert needs_ptf(row.ptf_required), (prod, row.ptf_required)


def test_summary_counts_ptf_rows():
    s = DB.summarize(DB.df, "3.2")
    assert s["ptf"] == 29 and s["yes"] == 153 and s["no"] == 44


def test_matches_reference_workbook_exactly():
    ws = load_workbook(REF)["Filtered"]
    key = {(r.vendor, r["product"], r.vrm): r for _, r in DB.df.iterrows()}
    for row in range(2, ws.max_row + 1):
        v, p, vrm, st, ptf = (str(ws.cell(row, c).value) for c in range(1, 6))
        r = key[(v, p, vrm)]
        assert r.zos32 == st and r.ptf_required == ptf, (p, vrm)


def test_excel_has_yellow_ptf_cells_like_reference():
    xl = build_excel({"All": DB.df}, version="3.2")
    path = os.path.join(tempfile.mkdtemp(), "o.xlsx")
    open(path, "wb").write(xl)
    ours, ref = load_workbook(path)["All"], load_workbook(REF)["Filtered"]
    assert ours.max_row == ref.max_row
    for row in range(2, ref.max_row + 1):
        f_ref = ref.cell(row, 5).fill.fgColor.rgb if ref.cell(row, 5).fill.fill_type else None
        f_our = ours.cell(row, 5).fill.fgColor.rgb if ours.cell(row, 5).fill.fill_type else None
        assert (f_ref or "").endswith(PTF_FILL) == (f_our or "").endswith(PTF_FILL), row
        assert ours.cell(row, 5).value == ref.cell(row, 5).value, row
    yellow = sum(1 for r in range(2, ours.max_row + 1) if ours.cell(r, 5).fill.fill_type)
    assert yellow == 29
    assert "PTF Worklist" in load_workbook(path).sheetnames


def test_worklist():
    wl = ptf_worklist(DB.df, "3.2")
    assert len(wl) == 29
    ids = unique_ptf_ids(DB.df, "3.2")
    assert "LU16693" in ids and "LU05158" in ids and "LU18249" not in ids and "UJ98216" not in ids


def test_filter_ptf_required():
    got = DB.filter(ptf="PTF required", version="3.2")
    assert len(got) == 29 and got["ptf_required"].map(needs_ptf).all()


def test_importer_is_idempotent():
    tmp = tempfile.mkdtemp()
    shutil.copytree(os.path.join(ROOT, "data"), os.path.join(tmp, "data"))
    assert import_xlsx.main([REF, "--data-dir", os.path.join(tmp, "data")]) == 0
    before = {f: open(f, "rb").read() for f in glob.glob(os.path.join(tmp, "data", "*.csv"))}
    import_xlsx.main([REF, "--data-dir", os.path.join(tmp, "data")])
    assert before == {f: open(f, "rb").read() for f in before}


def test_streamlit_app_runs_without_errors():
    try:
        from streamlit.testing.v1 import AppTest
    except Exception:
        return
    at = AppTest.from_file(os.path.join(ROOT, "app.py"), default_timeout=60).run()
    assert not at.exception, at.exception
    at.selectbox(key="zos_ver").select("3.1").run()
    assert not at.exception, at.exception
    at.selectbox(key="zos_ver").select("2.5").run()
    assert not at.exception, at.exception


if __name__ == "__main__":
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_"):
            try:
                fn(); print("PASS", name)
            except Exception as e:
                fails += 1; print("FAIL", name, "->", repr(e)[:300])
    sys.exit(1 if fails else 0)
